from io import BytesIO
from fpdf import FPDF
from docx import Document

def make_txt(content):
    return content.encode("utf-8")

def make_docx(content):
    doc = Document()
    for block in content.split("\n"):
        if block.strip():
            doc.add_paragraph(block)
        else:
            doc.add_paragraph("")
    stream = BytesIO()
    doc.save(stream)
    stream.seek(0)
    return stream.getvalue()

def make_pdf(content):
    pdf = FPDF()
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.add_page()
    pdf.set_font("Arial", size=11)

    for line in content.splitlines():
        safe = line.encode("latin-1", "replace").decode("latin-1")
        if not safe.strip():
            pdf.ln(4)
        else:
            pdf.multi_cell(0, 7, safe)

    return bytes(pdf.output(dest="S"))
